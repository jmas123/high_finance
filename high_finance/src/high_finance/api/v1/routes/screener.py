"""
Stock Screener – API Routes
===========================
Endpoints:
    GET  /api/v1/screener           - Run a screen with filters
    GET  /api/v1/screener/export/excel - Export results as .xlsx
    POST /api/v1/screener/saved     - Save a screen
    GET  /api/v1/screener/saved     - List saved screens
    DELETE /api/v1/screener/saved/{id} - Delete a saved screen
"""

import logging
from datetime import date
from io import BytesIO
from typing import Optional

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from high_finance.db.session import get_db
from high_finance.services import screener as svc
from high_finance.core.security import get_current_user, get_current_user_optional, User
from high_finance.api.v1.schemas.screener import (
    ScreenerResponse,
    ScreenerResult,
    SavedScreenCreate,
    SavedScreenOut,
)

logger = logging.getLogger(__name__)

router = APIRouter()


# ═══════════════════════════════════════════════════════════════════════════════
# SCREENER
# ═══════════════════════════════════════════════════════════════════════════════

@router.get("", response_model=ScreenerResponse)
async def run_screener(
    # Size
    market_cap_min: Optional[int] = Query(None, description="Min market cap in millions"),
    market_cap_max: Optional[int] = Query(None, description="Max market cap in millions"),
    # Valuation
    pe_ratio_min: Optional[float] = Query(None),
    pe_ratio_max: Optional[float] = Query(None),
    forward_pe_min: Optional[float] = Query(None),
    forward_pe_max: Optional[float] = Query(None),
    pb_ratio_min: Optional[float] = Query(None),
    pb_ratio_max: Optional[float] = Query(None),
    ev_to_ebitda_min: Optional[float] = Query(None),
    ev_to_ebitda_max: Optional[float] = Query(None),
    # Quality & Profitability
    roe_min: Optional[float] = Query(None, description="Decimal, e.g., 0.15 = 15%"),
    roe_max: Optional[float] = Query(None),
    profit_margin_min: Optional[float] = Query(None, description="Decimal, e.g., 0.1 = 10%"),
    profit_margin_max: Optional[float] = Query(None),
    debt_to_equity_min: Optional[float] = Query(None),
    debt_to_equity_max: Optional[float] = Query(None),
    fcf_yield_min: Optional[float] = Query(None, description="Decimal, e.g., 0.05 = 5%"),
    fcf_yield_max: Optional[float] = Query(None),
    beta_min: Optional[float] = Query(None),
    beta_max: Optional[float] = Query(None),
    # Growth & Dividend
    revenue_growth_min: Optional[float] = Query(None, description="Decimal, e.g., 0.1 = 10%"),
    revenue_growth_max: Optional[float] = Query(None),
    dividend_yield_min: Optional[float] = Query(None),
    dividend_yield_max: Optional[float] = Query(None),
    # Price Technicals
    price_vs_200ma_min: Optional[float] = Query(None, description="Decimal, e.g., 0.05 = 5% above"),
    price_vs_200ma_max: Optional[float] = Query(None),
    price_vs_52w_high_min: Optional[float] = Query(None, description="Decimal, e.g., -0.1 = within 10%"),
    price_vs_52w_high_max: Optional[float] = Query(None),
    # Sectors
    sectors: Optional[str] = Query(None, description="Comma-separated sector names"),
    # Insider
    insider_net_shares_90d_min: Optional[int] = Query(None),
    has_insider_buys_90d: Optional[bool] = Query(None),
    # Congress
    congress_net_value_90d_min: Optional[int] = Query(None),
    # Sort & Pagination
    sort_by: str = Query("market_cap", description="Field to sort by"),
    sort_order: str = Query("desc", description="asc or desc"),
    limit: int = Query(50, ge=1, le=200, description="Results per page (max 200)"),
    offset: int = Query(0, ge=0),
    # DB
    db: Session = Depends(get_db),
):
    """
    Run a stock screener query with the specified filters.

    Filters are combined with AND logic. Null values are excluded by default
    when a min/max filter is applied for that field.
    """
    # Collect filters applied for response
    filters_applied = {}
    local_vars = locals()
    filter_params = [
        "market_cap_min", "market_cap_max", "pe_ratio_min", "pe_ratio_max",
        "forward_pe_min", "forward_pe_max", "pb_ratio_min", "pb_ratio_max",
        "ev_to_ebitda_min", "ev_to_ebitda_max",
        "roe_min", "roe_max", "profit_margin_min", "profit_margin_max",
        "debt_to_equity_min", "debt_to_equity_max", "fcf_yield_min", "fcf_yield_max",
        "beta_min", "beta_max",
        "revenue_growth_min", "revenue_growth_max",
        "dividend_yield_min", "dividend_yield_max", "price_vs_200ma_min", "price_vs_200ma_max",
        "price_vs_52w_high_min", "price_vs_52w_high_max", "sectors",
        "insider_net_shares_90d_min", "has_insider_buys_90d", "congress_net_value_90d_min",
        "sort_by", "sort_order", "limit", "offset",
    ]
    for param in filter_params:
        val = local_vars.get(param)
        if val is not None:
            filters_applied[param] = val

    results, total_count, cached_at = await svc.run_screen(
        db,
        market_cap_min=market_cap_min,
        market_cap_max=market_cap_max,
        pe_ratio_min=pe_ratio_min,
        pe_ratio_max=pe_ratio_max,
        forward_pe_min=forward_pe_min,
        forward_pe_max=forward_pe_max,
        pb_ratio_min=pb_ratio_min,
        pb_ratio_max=pb_ratio_max,
        ev_to_ebitda_min=ev_to_ebitda_min,
        ev_to_ebitda_max=ev_to_ebitda_max,
        roe_min=roe_min,
        roe_max=roe_max,
        profit_margin_min=profit_margin_min,
        profit_margin_max=profit_margin_max,
        debt_to_equity_min=debt_to_equity_min,
        debt_to_equity_max=debt_to_equity_max,
        fcf_yield_min=fcf_yield_min,
        fcf_yield_max=fcf_yield_max,
        beta_min=beta_min,
        beta_max=beta_max,
        revenue_growth_min=revenue_growth_min,
        revenue_growth_max=revenue_growth_max,
        dividend_yield_min=dividend_yield_min,
        dividend_yield_max=dividend_yield_max,
        price_vs_200ma_min=price_vs_200ma_min,
        price_vs_200ma_max=price_vs_200ma_max,
        price_vs_52w_high_min=price_vs_52w_high_min,
        price_vs_52w_high_max=price_vs_52w_high_max,
        sectors=sectors,
        insider_net_shares_90d_min=insider_net_shares_90d_min,
        has_insider_buys_90d=has_insider_buys_90d,
        congress_net_value_90d_min=congress_net_value_90d_min,
        sort_by=sort_by,
        sort_order=sort_order,
        limit=limit,
        offset=offset,
    )

    universe_total = await svc.get_universe_total(db)

    return ScreenerResponse(
        results=[ScreenerResult(**r) for r in results],
        total_count=total_count,
        universe_total=universe_total,
        filters_applied=filters_applied,
        cached_at=cached_at,
    )


@router.get("/count")
async def get_screener_count(
    # Same filters as main screener endpoint
    market_cap_min: Optional[int] = Query(None),
    market_cap_max: Optional[int] = Query(None),
    pe_ratio_min: Optional[float] = Query(None),
    pe_ratio_max: Optional[float] = Query(None),
    forward_pe_min: Optional[float] = Query(None),
    forward_pe_max: Optional[float] = Query(None),
    pb_ratio_min: Optional[float] = Query(None),
    pb_ratio_max: Optional[float] = Query(None),
    ev_to_ebitda_min: Optional[float] = Query(None),
    ev_to_ebitda_max: Optional[float] = Query(None),
    roe_min: Optional[float] = Query(None),
    roe_max: Optional[float] = Query(None),
    profit_margin_min: Optional[float] = Query(None),
    profit_margin_max: Optional[float] = Query(None),
    debt_to_equity_min: Optional[float] = Query(None),
    debt_to_equity_max: Optional[float] = Query(None),
    fcf_yield_min: Optional[float] = Query(None),
    fcf_yield_max: Optional[float] = Query(None),
    beta_min: Optional[float] = Query(None),
    beta_max: Optional[float] = Query(None),
    revenue_growth_min: Optional[float] = Query(None),
    revenue_growth_max: Optional[float] = Query(None),
    dividend_yield_min: Optional[float] = Query(None),
    dividend_yield_max: Optional[float] = Query(None),
    price_vs_200ma_min: Optional[float] = Query(None),
    price_vs_200ma_max: Optional[float] = Query(None),
    price_vs_52w_high_min: Optional[float] = Query(None),
    price_vs_52w_high_max: Optional[float] = Query(None),
    sectors: Optional[str] = Query(None),
    insider_net_shares_90d_min: Optional[int] = Query(None),
    has_insider_buys_90d: Optional[bool] = Query(None),
    congress_net_value_90d_min: Optional[int] = Query(None),
    db: Session = Depends(get_db),
):
    """Return just the count matching current filters. Fast — no data returned."""
    _, total_count, _ = await svc.run_screen(
        db,
        market_cap_min=market_cap_min, market_cap_max=market_cap_max,
        pe_ratio_min=pe_ratio_min, pe_ratio_max=pe_ratio_max,
        forward_pe_min=forward_pe_min, forward_pe_max=forward_pe_max,
        pb_ratio_min=pb_ratio_min, pb_ratio_max=pb_ratio_max,
        ev_to_ebitda_min=ev_to_ebitda_min, ev_to_ebitda_max=ev_to_ebitda_max,
        roe_min=roe_min, roe_max=roe_max,
        profit_margin_min=profit_margin_min, profit_margin_max=profit_margin_max,
        debt_to_equity_min=debt_to_equity_min, debt_to_equity_max=debt_to_equity_max,
        fcf_yield_min=fcf_yield_min, fcf_yield_max=fcf_yield_max,
        beta_min=beta_min, beta_max=beta_max,
        revenue_growth_min=revenue_growth_min, revenue_growth_max=revenue_growth_max,
        dividend_yield_min=dividend_yield_min, dividend_yield_max=dividend_yield_max,
        price_vs_200ma_min=price_vs_200ma_min, price_vs_200ma_max=price_vs_200ma_max,
        price_vs_52w_high_min=price_vs_52w_high_min, price_vs_52w_high_max=price_vs_52w_high_max,
        sectors=sectors,
        insider_net_shares_90d_min=insider_net_shares_90d_min,
        has_insider_buys_90d=has_insider_buys_90d,
        congress_net_value_90d_min=congress_net_value_90d_min,
        limit=0, offset=0,
    )
    universe_total = await svc.get_universe_total(db)
    return {"count": total_count, "universe_total": universe_total}


@router.get("/sectors")
async def get_sector_counts(db: Session = Depends(get_db)):
    """Get count of stocks per sector in the universe."""
    return await svc.get_sector_counts(db)


# ═══════════════════════════════════════════════════════════════════════════════
# EXCEL EXPORT
# ═══════════════════════════════════════════════════════════════════════════════

EXCEL_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

# Column definitions for the Excel export
_EXPORT_COLUMNS = [
    ("Ticker",         "ticker",               "text"),
    ("Company",        "company_name",          "text"),
    ("Sector",         "sector",                "text"),
    ("Industry",       "industry",              "text"),
    ("Market Cap",     "market_cap",            "currency_int"),
    ("Price",          "current_price",         "currency"),
    ("1D Change",      "price_change_1d",       "pct"),
    ("P/E",            "pe_ratio",              "num"),
    ("Forward P/E",    "forward_pe",            "num"),
    ("P/B",            "pb_ratio",              "num"),
    ("EV/EBITDA",      "ev_to_ebitda",          "num"),
    ("ROE",            "roe",                   "pct"),
    ("Profit Margin",  "profit_margin",         "pct"),
    ("D/E",            "debt_to_equity",        "num"),
    ("FCF Yield",      "fcf_yield",             "pct"),
    ("Beta",           "beta",                  "num"),
    ("Rev Growth",     "revenue_growth",        "pct"),
    ("Div Yield",      "dividend_yield",        "pct"),
    ("Price vs 200MA", "price_vs_200ma",        "pct"),
    ("Price vs 52W High", "price_vs_52w_high",  "pct"),
    ("Insider Net Shares (90d)", "insider_net_shares_90d", "int"),
    ("Insider Buys (90d)",       "insider_buys_90d",       "int"),
    ("Insider Sells (90d)",      "insider_sells_90d",      "int"),
    ("Congress Net Value (90d)", "congress_net_value_90d",  "int"),
    ("Congress Buys (90d)",      "congress_buys_90d",       "int"),
    ("Congress Sells (90d)",     "congress_sells_90d",      "int"),
]


def _build_screener_workbook(rows: list[dict]) -> BytesIO:
    """Build an openpyxl workbook from screener result dicts and return as BytesIO."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Screener Results"

    # --- Styles ---
    header_fill = PatternFill(start_color="2C3E50", end_color="2C3E50", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=11)
    thin_border = Border(
        left=Side(style="thin", color="CCCCCC"),
        right=Side(style="thin", color="CCCCCC"),
        top=Side(style="thin", color="CCCCCC"),
        bottom=Side(style="thin", color="CCCCCC"),
    )
    pct_fmt = '0.00%'
    currency_fmt = '$#,##0.00'
    currency_int_fmt = '$#,##0'
    num_fmt = '0.00'
    int_fmt = '#,##0'

    # --- Header row ---
    for col_idx, (label, _, _) in enumerate(_EXPORT_COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=label)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
        cell.border = thin_border

    # --- Data rows ---
    for row_idx, row_data in enumerate(rows, start=2):
        for col_idx, (_, key, col_type) in enumerate(_EXPORT_COLUMNS, start=1):
            value = row_data.get(key)
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.border = thin_border

            if value is None:
                cell.value = None
                continue

            cell.value = value
            if col_type == "pct":
                cell.number_format = pct_fmt
                cell.alignment = Alignment(horizontal="right")
            elif col_type == "currency":
                cell.number_format = currency_fmt
                cell.alignment = Alignment(horizontal="right")
            elif col_type == "currency_int":
                cell.number_format = currency_int_fmt
                cell.alignment = Alignment(horizontal="right")
            elif col_type == "num":
                cell.number_format = num_fmt
                cell.alignment = Alignment(horizontal="right")
            elif col_type == "int":
                cell.number_format = int_fmt
                cell.alignment = Alignment(horizontal="right")

    # --- Auto-size columns ---
    for col_idx, (label, _, _) in enumerate(_EXPORT_COLUMNS, start=1):
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = max(len(label) + 4, 12)

    # Freeze header row
    ws.freeze_panes = "A2"

    # Auto-filter
    ws.auto_filter.ref = f"A1:{get_column_letter(len(_EXPORT_COLUMNS))}1"

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


@router.get("/export/excel")
async def export_screener_excel(
    # Same filters as main screener endpoint
    market_cap_min: Optional[int] = Query(None),
    market_cap_max: Optional[int] = Query(None),
    pe_ratio_min: Optional[float] = Query(None),
    pe_ratio_max: Optional[float] = Query(None),
    forward_pe_min: Optional[float] = Query(None),
    forward_pe_max: Optional[float] = Query(None),
    pb_ratio_min: Optional[float] = Query(None),
    pb_ratio_max: Optional[float] = Query(None),
    ev_to_ebitda_min: Optional[float] = Query(None),
    ev_to_ebitda_max: Optional[float] = Query(None),
    roe_min: Optional[float] = Query(None),
    roe_max: Optional[float] = Query(None),
    profit_margin_min: Optional[float] = Query(None),
    profit_margin_max: Optional[float] = Query(None),
    debt_to_equity_min: Optional[float] = Query(None),
    debt_to_equity_max: Optional[float] = Query(None),
    fcf_yield_min: Optional[float] = Query(None),
    fcf_yield_max: Optional[float] = Query(None),
    beta_min: Optional[float] = Query(None),
    beta_max: Optional[float] = Query(None),
    revenue_growth_min: Optional[float] = Query(None),
    revenue_growth_max: Optional[float] = Query(None),
    dividend_yield_min: Optional[float] = Query(None),
    dividend_yield_max: Optional[float] = Query(None),
    price_vs_200ma_min: Optional[float] = Query(None),
    price_vs_200ma_max: Optional[float] = Query(None),
    price_vs_52w_high_min: Optional[float] = Query(None),
    price_vs_52w_high_max: Optional[float] = Query(None),
    sectors: Optional[str] = Query(None),
    insider_net_shares_90d_min: Optional[int] = Query(None),
    has_insider_buys_90d: Optional[bool] = Query(None),
    congress_net_value_90d_min: Optional[int] = Query(None),
    sort_by: str = Query("market_cap"),
    sort_order: str = Query("desc"),
    db: Session = Depends(get_db),
):
    """
    Export all matching screener results as a formatted .xlsx file.

    Uses the same filter params as the main screener endpoint, but
    returns up to 10,000 rows (no pagination) as an Excel download.
    """
    try:
        results, total_count, _ = await svc.run_screen(
            db,
            market_cap_min=market_cap_min, market_cap_max=market_cap_max,
            pe_ratio_min=pe_ratio_min, pe_ratio_max=pe_ratio_max,
            forward_pe_min=forward_pe_min, forward_pe_max=forward_pe_max,
            pb_ratio_min=pb_ratio_min, pb_ratio_max=pb_ratio_max,
            ev_to_ebitda_min=ev_to_ebitda_min, ev_to_ebitda_max=ev_to_ebitda_max,
            roe_min=roe_min, roe_max=roe_max,
            profit_margin_min=profit_margin_min, profit_margin_max=profit_margin_max,
            debt_to_equity_min=debt_to_equity_min, debt_to_equity_max=debt_to_equity_max,
            fcf_yield_min=fcf_yield_min, fcf_yield_max=fcf_yield_max,
            beta_min=beta_min, beta_max=beta_max,
            revenue_growth_min=revenue_growth_min, revenue_growth_max=revenue_growth_max,
            dividend_yield_min=dividend_yield_min, dividend_yield_max=dividend_yield_max,
            price_vs_200ma_min=price_vs_200ma_min, price_vs_200ma_max=price_vs_200ma_max,
            price_vs_52w_high_min=price_vs_52w_high_min, price_vs_52w_high_max=price_vs_52w_high_max,
            sectors=sectors,
            insider_net_shares_90d_min=insider_net_shares_90d_min,
            has_insider_buys_90d=has_insider_buys_90d,
            congress_net_value_90d_min=congress_net_value_90d_min,
            sort_by=sort_by,
            sort_order=sort_order,
            limit=10000,
            offset=0,
        )
    except Exception as e:
        logger.exception(f"Screener export failed: {e}")
        raise HTTPException(422, "Unable to generate export. Please try again.")

    if not results:
        raise HTTPException(404, "No stocks match your filters.")

    excel_bytes = _build_screener_workbook(results)
    filename = f"screener_{date.today().isoformat()}.xlsx"

    return StreamingResponse(
        excel_bytes,
        media_type=EXCEL_MIME,
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


# ═══════════════════════════════════════════════════════════════════════════════
# SAVED SCREENS
# ═══════════════════════════════════════════════════════════════════════════════

@router.get("/saved", response_model=list[SavedScreenOut])
async def list_saved_screens(
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
):
    """List all saved screens for the current user."""
    screens = await svc.list_saved_screens(db, user.id)
    return screens


@router.post("/saved", response_model=SavedScreenOut, status_code=201)
async def create_saved_screen(
    body: SavedScreenCreate,
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
):
    """Save a new screen configuration."""
    screen = await svc.create_saved_screen(
        db,
        owner_id=user.id,
        name=body.name,
        filters=body.filters,
    )
    return screen


@router.delete("/saved/{screen_id}", status_code=204)
async def delete_saved_screen(
    screen_id: int,
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
):
    """Delete a saved screen."""
    deleted = await svc.delete_saved_screen(db, screen_id, user.id)
    if not deleted:
        raise HTTPException(404, "Saved screen not found")
